from pathlib import Path
from docx import Document
from openpyxl import load_workbook

ROOT = Path(r"C:\Users\cveda\Acads\Sem-5\Software engineering\Labs\Fusion_2nd_model")
SOURCE = ROOT / "ELM_Specifications_Bundle_v1.4_FINAL"
OUTPUT = ROOT / ".elm_context"
OUTPUT.mkdir(exist_ok=True)

def extract_docx(path: Path) -> str:
    doc = Document(path)
    sections = []

    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text:
            sections.append(text)

    for table_index, table in enumerate(doc.tables, start=1):
        sections.append(f"\n[TABLE {table_index}]")
        for row in table.rows:
            sections.append(" | ".join(cell.text.strip() for cell in row.cells))

    return "\n".join(sections)

def extract_xlsx(path: Path) -> str:
    workbook = load_workbook(path, data_only=False)
    sections = []

    for sheet in workbook.worksheets:
        sections.append(f"\n[WORKSHEET: {sheet.title}]")

        for row in sheet.iter_rows(values_only=True):
            values = ["" if value is None else str(value) for value in row]

            if any(value.strip() for value in values):
                sections.append(" | ".join(values))

    return "\n".join(sections)

# --- NEW DEBUG SECTION ---

if not SOURCE.exists():
    print(f"ERROR: The source folder does not exist at:\n{SOURCE}")
else:
    print(f" Scanning directory: {SOURCE}")
    
    # Check if folder is completely empty
    files = list(SOURCE.iterdir())
    if not files:
        print("Folder is completely empty!")
        
    for path in sorted(files):
        print(f"Found file: {path.name} (Suffix: {path.suffix.lower()})")
        
        if path.suffix.lower() == ".docx":
            content = extract_docx(path)
        elif path.suffix.lower() == ".xlsx":
            content = extract_xlsx(path)
        else:
            print(f"   -> Skipping {path.name} - Not a .docx or .xlsx file")
            continue

        output_file = OUTPUT / f"{path.stem}.txt"
        output_file.write_text(content, encoding="utf-8")
        print(f"   -> Created {output_file}")